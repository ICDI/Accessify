#!/usr/bin/env python3
"""Gera o formulário do Estudo 2 (docs/pesquisa/estudo2_formulario.xlsx) e, opcionalmente, os
cartões de tarefa em Markdown usados em docs/pesquisa/estudo2_tarefas.md, a partir de uma única
fonte (lista T abaixo). Uso:
   python3 eval/tools/estudo2_gerar.py [--xlsx SAIDA.xlsx] [--cartoes SAIDA.md] [--trocar ID=RESERVA ...]
--trocar põe uma reserva no lugar de uma tarefa (mesmo grupo e mesma posição na sessão), por
exemplo --trocar CSBC-2=IHC-3 quando o piloto mostrar que o NVDA não lê o endereço de links sem
nome. Requer openpyxl. Depois de gerar, recalcular as fórmulas (LibreOffice) antes de abrir em
visualizadores que só leem valores em cache.
Protocolo: docs/pesquisa/estudo2_tarefas.md. Conferência das respostas: eval/tools/estudo2_conferir.py."""
import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit('requer openpyxl (pip install openpyxl)')

REPO = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..')


def opcao(nome, padrao=None):
    return sys.argv[sys.argv.index(nome) + 1] if nome in sys.argv else padrao


SAIDA_XLSX = opcao('--xlsx', os.path.join(REPO, 'docs/pesquisa/estudo2_formulario.xlsx'))
SAIDA_MD = opcao('--cartoes')
TROCAS = dict(a.split('=') for a in sys.argv[sys.argv.index('--trocar') + 1:] if '=' in a) if '--trocar' in sys.argv else {}

PAINEL_S_POR_PAGINA = 20  # operação do painel por página na condição C (Alt+Shift+A, ouvir, Aplicar, Esc)

REGRA = 'Regra geral: S = resposta correta, sem ajuda, em até 300 s; F = resposta errada, desistência, ajuda para avançar ou 300 s.'

T = [
    # ---------------- Núcleo ----------------
    dict(id='SOL-1', conj='Núcleo', grupo='G1', site='SOL — busca integrada', url='https://sol.sbc.org.br/busca/',
         titulo='Buscar um artigo na SOL e dizer o primeiro autor', tipo='baixo', pag=2, est_s=105, extra=0,
         instr_a='Na busca da biblioteca digital da SBC, a SOL, procure nos Anais de Eventos o artigo "Can AI Judge Usability? A Comparative Analysis of Generative Tools on Climate Conference Websites" e diga o nome do primeiro autor.',
         saida_a='Danilo Teixeira Lima (no resultado deve aparecer como "Lima, Danilo Teixeira"; confirmar no piloto)',
         crit_a='S: "Danilo Teixeira Lima", "Danilo Lima" ou "Lima". P: chega ao resultado certo, mas diz outro autor do artigo. F: não consegue enviar a busca (a base "Anais de Eventos" é obrigatória e o aviso é só visual), outro nome, desistência.',
         instr_b='Na busca da biblioteca digital da SBC, a SOL, procure nos Anais de Eventos o artigo "Teaching or Manipulating? On the Adoption of Bright and Deceptive Patterns by Duolingo" e diga o nome do primeiro autor.',
         saida_b='João Vítor Castro (provavelmente "Castro, João Vítor"; confirmar no piloto)',
         crit_b='S: "João Vítor Castro", "João Castro" ou "Castro". P: chega ao resultado certo, mas diz o segundo autor (George Valença). F: não envia a busca, outro nome, desistência.',
         caminho='Busca integrada: marcar a caixa "Anais de Eventos" → campo "Título" → digitar parte do título e Enter → na página de resultados, ler a linha de autores do item.',
         efeito='Baixo. O campo "Título" já tem rótulo. A extensão nomeia o campo sem rótulo e os selects de data pelo atributo name ("Query", "Field 7 from Month") e acrescenta "(abre em nova aba)" aos links de resultado; não trata a base obrigatória com aviso só visual, que é a barreira real.',
         riscos='As respostas foram conferidas no sumário do IHC 2025 (cópia do Wayback de 22/10/2026), não no resultado da busca ao vivo: o servidor sol.sbc.org.br recusa conexões desta rede desde a exploração. Conferir no piloto que a busca pelo título traz o artigo no topo. Registrar se ela tentou enviar sem marcar a base e se abriu o artigo na nova aba. Espaçar as tarefas da SOL; reserva: SOL-4 (servidor books-sol).'),
    dict(id='SOL-2', conj='Núcleo', grupo='G2', site='SOL — anais do IHC 2025', url='https://sol.sbc.org.br/index.php/ihc/issue/view/1582',
         titulo='Dizer em que trilha dos anais do IHC 2025 está um artigo', tipo='controle', pag=1, est_s=45, extra=0,
         instr_a='Nos anais do IHC 2025, na SOL, diga em qual trilha foi publicado o artigo "Falando, desenhando e avaliando: workshops para a construção de um repositório sobre tecnologias para pessoas idosas".',
         saida_a='Relatos de Experiência',
         crit_a='S: "Relatos de Experiência" (com ou sem "Trilha"). P: chega ao artigo, mas não identifica a trilha. F: outra trilha (Abertura, Ideias Inovadoras, Prêmio, Artigos de Pesquisa), desistência.',
         instr_b='Nos anais do IHC 2025, na SOL, diga em qual trilha foi publicado o artigo "Diagnóstico da Acessibilidade Digital em Serviços Públicos Prioritários usando a Lei de Acesso à Informação".',
         saida_b='Ideias Inovadoras e Resultados Emergentes em IHC',
         crit_b='S: "Ideias Inovadoras e Resultados Emergentes em IHC" (aceitar "Ideias Inovadoras"). P: chega ao artigo, mas não identifica a trilha. F: outra trilha, desistência.',
         caminho='Busca do NVDA por parte do título → Shift+H até o cabeçalho de nível 2 da trilha (ou abrir o artigo e ler "Você está aqui").',
         efeito='Nenhum (controle). Trilhas são h2 e títulos são links com texto. A extensão só muda contraste e expõe um link "Buscar" oculto, fora do caminho.',
         riscos='Mesmo servidor da SOL-1. Posição absoluta diferente (A é o 4º item da página, B o 21º): registrar a estratégia (busca de texto, lista de links, leitura linear).'),
    dict(id='CAPES-1', conj='Núcleo', grupo='G1', site='Portal de Periódicos CAPES — Lista A–Z', url='https://www.periodicos.capes.gov.br/index.php/acervo/lista-a-z-periodicos.html',
         titulo='Achar o ISSN de um periódico na Lista A–Z', tipo='baixo', pag=3, est_s=90, extra=0,
         instr_a='No Portal de Periódicos da CAPES, descubra qual é o ISSN do periódico Journal of the Brazilian Computer Society.',
         saida_a='0104-6500',
         crit_a='S: 0104-6500 (hífen opcional). P: chegou à página do periódico sem ler o ISSN, ou deu um ISSN lido fora do Portal (SciELO). F: ISSN de outro título, desistência.',
         instr_b='No Portal de Periódicos da CAPES, descubra qual é o ISSN do periódico Journal of the Brazilian Chemical Society.',
         saida_b='0103-5053',
         crit_b='S: 0103-5053 (hífen opcional). P: chegou à página do periódico sem ler o ISSN, ou deu um ISSN lido fora do Portal (SciELO). F: ISSN de outro título, desistência.',
         caminho='Campo de busca da lista (anunciado "Olá.") → digitar o título e Enter → link do resultado único → tabela de detalhes, linha "ISSN".',
         efeito='Baixo (nulo a levemente negativo). É a página com a maior queda do axe no Estudo 1 (172 → 40), mas o caminho quase não muda: o rótulo enganoso "Olá." do campo continua e a tabela de detalhes já é acessível. A extensão renomeia botões que tinham bom nome ("Exibir lista", "Página seguinte") para "Acionar botão N", fora do caminho.',
         riscos='O cabeçalho do detalhe e o botão "Acesse" levam ao SciELO (outro site, outro ISSN). Fechar o banner de cookies antes de cronometrar, nas duas condições.'),
    dict(id='CAPES-2', conj='Núcleo', grupo='G2', site='Portal de Periódicos CAPES — Suporte regional', url='https://www.periodicos.capes.gov.br/index.php/ajuda/suporte-regional.html',
         titulo='E-mail do suporte regional (mapa de botões sem nome)', tipo='baixo', pag=1, est_s=110, extra=0,
         instr_a='No Portal de Periódicos da CAPES, encontre o e-mail de contato do suporte regional para a região Centro-Oeste.',
         saida_a='lucianasetubal@unb.br ou seanio.avelino@unb.br (qualquer um)',
         crit_a='S: um dos dois e-mails @unb.br. P: abre a região Centro-Oeste, mas dá só o nome ou o telefone, ou soletra o e-mail com 1 caractere errado. F: e-mail de outra região, desistência.',
         instr_b='No Portal de Periódicos da CAPES, encontre o e-mail de contato do suporte regional para a região Sudeste.',
         saida_b='cref@sibi.ufrj.br',
         crit_b='S: cref@sibi.ufrj.br. P: abre a região Sudeste, mas dá só o nome de uma bibliotecária, ou soletra o e-mail com 1 caractere errado. F: e-mail de outra região, desistência.',
         caminho='Cabeçalho "Atendimento por região" → percorrer os 5 botões do mapa abrindo cada menu até ouvir a região pedida → ler nome e e-mail.',
         efeito='Baixo. Os 5 botões do mapa não têm nome (button-name crítico); a extensão lhes dá "Acionar botão N", sem a região, que continua só dentro do menu fechado. Mede o defeito D-7 numa barreira que a extensão diz corrigir.',
         riscos='Janela maximizada (≥ 1280 px) e zoom 100%: abaixo de 768 px CSS aparece uma lista alternativa e a barreira some. E-mails pessoais podem mudar: rodar o conferir na véspera. Anotar quantos menus foram abertos. A variante A (Centro-Oeste, 3º botão) fica antes da B (Sudeste, 4º) no mapa e é feita com a extensão, para não somar a vantagem de posição à da aprendizagem, que favorece S neste grupo.'),
    dict(id='IHC-1', conj='Núcleo', grupo='G2', site='IHC 2026 — Sessões Técnicas', url='https://ihc.sbc.org.br/2026/',
         titulo='Dia, horário e sala em que um trabalho foi apresentado', tipo='baixo', pag=3, est_s=105, extra=0,
         instr_a='No site do IHC 2026, sem baixar arquivos, descubra em que dia, horário e sala foi apresentado o artigo "Explorando lacunas e oportunidades em aplicativos móveis para suporte ao luto com aplicabilidade profissional".',
         saida_a='Quinta-feira, 22/10, das 08h30 às 10h00, sala Algaroba (Sessão Técnica 10)',
         crit_a='S: dia (quinta ou 22/10) + início 08h30 + sala Algaroba. P: 2 dos 3 dados, ou só a sessão certa. F: dados de outra sessão, desistência.',
         instr_b='No site do IHC 2026, sem baixar arquivos, descubra em que dia, horário e sala foi apresentado o artigo "Lições Aprendidas sobre a Comunicabilidade da Ferramenta FAKECHECK: Uma Avaliação da Plataforma Web para Detecção de Notícias Falsas".',
         saida_b='Quarta-feira, 21/10, das 14h às 15h30, sala Juazeiro (Sessão Técnica 7)',
         crit_b='S: dia (quarta ou 21/10) + início 14h + sala Juazeiro. P: 2 dos 3 dados, ou só a sessão certa. F: dados de outra sessão, desistência.',
         caminho='Início → "Programação" → "Sessões Técnicas" → busca do NVDA por uma palavra do título → Shift+H até o tema da sessão → lista Data/Horário/Coordenador/Sala.',
         efeito='Baixo, provavelmente negativo. Os links do caminho já têm texto. A extensão remove o aria-hidden de cerca de 120 ícones decorativos da página Sessões Técnicas, que passam a ser lidos como "gráfico" justamente na lista Data/Horário/Sala (defeito D-6).',
         riscos='A Programação oferece "Programação Acessível" em DOCX/PDF: a instrução diz "sem baixar arquivos". Títulos longos: permitir repetir a instrução e contar as repetições. Os alvos são o 3º de 4 trabalhos em sessões de mesma estrutura.'),
    dict(id='IHC-2', conj='Núcleo', grupo='G1', site='IHC 2026 — Programação Geral', url='https://ihc.sbc.org.br/2026/index.php/programacao/',
         titulo='Qual sessão técnica aconteceu num dia e horário (abas + tabela)', tipo='controle', pag=1, est_s=70, extra=0,
         instr_a='Na página de programação do IHC 2026, sem baixar arquivos, descubra qual sessão técnica aconteceu na terça-feira, 20 de outubro, às 10h30.',
         saida_a='Sessão Técnica 2 — Diversidade e Justiça Social (Sala Juazeiro, 10h30–12h)',
         crit_a='S: número (2) ou tema (Diversidade e Justiça Social). P: lê a linha certa, mas não diz qual é a sessão técnica. F: outro dia (a aba de segunda abre por padrão), outra linha, desistência.',
         instr_b='Na página de programação do IHC 2026, sem baixar arquivos, descubra qual sessão técnica aconteceu na quarta-feira, 21 de outubro, às 10h30.',
         saida_b='Sessão Técnica 6 — Padrões Enganosos e Privacidade (Sala Juazeiro, 10h30–12h)',
         crit_b='S: número (6) ou tema (Padrões Enganosos e Privacidade). P: lê a linha certa, mas não diz qual é a sessão técnica. F: outro dia, outra linha, desistência.',
         caminho='Lista de abas → ativar a aba do dia → T até a tabela → linha "10h30 – 12h" → célula da sessão técnica.',
         efeito='Nenhum (controle). Abas ARIA e tabela com cabeçalho já acessíveis; a extensão não as altera. Estima o custo fixo de usar a extensão sem benefício possível.',
         riscos='A busca do NVDA não acha conteúdo de abas fechadas. Erro previsível de aba (segunda, aberta por padrão) é igual em A e B.'),
    dict(id='CSBC-2', conj='Núcleo', grupo='G1', site='CSBC 2026 — rodapé de patrocinadores', url='https://csbc.sbc.org.br/2026/',
         titulo='Qual empresa patrocina uma categoria (logo sem nome)', tipo='negativo', pag=1, est_s=50, extra=50,
         instr_a='No site do CSBC 2026, o Congresso da SBC, descubra qual empresa é a patrocinadora da categoria Cobre.',
         saida_a='Scherm (aceitar scherm.com.br, em qualquer pronúncia ou soletrado)',
         crit_a='S: "Scherm" ou o domínio. P: chega ao link da categoria Cobre, mas não consegue nomear a empresa. F: empresa de outra categoria, desistência.',
         instr_b='No site do CSBC 2026, o Congresso da SBC, descubra qual empresa é a patrocinadora da categoria Ouro.',
         saida_b='Qualcomm (aceitar o endereço qualcomm.com)',
         crit_b='S: "Qualcomm" ou o domínio. P: chega ao link da categoria Ouro, mas não consegue nomear a empresa. F: empresa de outra categoria, desistência.',
         caminho='Lista de cabeçalhos → "COBRE" (B: "OURO") → seta para baixo até o link do logo → ouvir o anúncio; se não identificar, NVDA+K ou seguir o link e ouvir o título do site.',
         efeito='Alto, NEGATIVO. O logo é um link com `<img alt="">` e sem nome; sem a extensão o NVDA tende a anunciar o endereço (scherm.com.br, qualcomm.com), o que basta. Com a extensão o link ganha nome genérico ("Abrir link 19" ou "Abrir link 11" pelo painel; o cabeçalho anterior, "COBRE" ou "OURO", no modo automático), que esconde o endereço. Mede o defeito D-7.',
         riscos='Só usar se o piloto confirmar que o NVDA da avaliadora anuncia o endereço de um link sem nome no Chrome; se não anunciar, trocar pela reserva IHC-3. Assimetria: Qualcomm é marca conhecida, Scherm não; por isso a variante mais difícil (A, Scherm) é feita sem a extensão (grupo G1), e tanto a variante quanto a aprendizagem jogam contra o efeito negativo previsto. Sites externos não recebem o Accessify.'),
    dict(id='EDU-1', conj='Núcleo', grupo='G2', site='EduComp 2026 — chamadas', url='https://educomp.sbc.org.br/2026/',
         titulo='Quem coordena uma chamada (links "Saiba mais")', tipo='positivo', pag=2, est_s=55, extra=0,
         instr_a='No site do EduComp 2026, descubra quem coordena a chamada de Minicursos.',
         saida_a='Maria Ivanilse Calderon Ribeiro (IFRO)',
         crit_a='S: o nome (ao menos prenome e um sobrenome; instituição opcional). P: abre a chamada certa, mas não lê o nome ou dá só o contato; ou responde os ministrantes listados em Programação > Minicursos. F: coordenação de outra chamada, desistência.',
         instr_b='No site do EduComp 2026, descubra quem coordena a chamada de Mesas temáticas.',
         saida_b='Jefferson Rodrigo Speck (Faculdade Donaduzzi/Unioeste)',
         crit_b='S: o nome (ao menos prenome e um sobrenome). P: abre a chamada certa, mas não lê o nome ou dá só o contato; ou responde os mediadores de uma mesa (Programação > Mesas temáticas). F: coordenação de outra chamada, desistência.',
         caminho='Cabeçalho do cartão da chamada → link "Saiba mais" (ou lista de links) → página da chamada, linha "Coordenação:" logo após o h1.',
         efeito='Baixo, POSITIVO. Os 7 links "Saiba mais" iguais passam a "Saiba mais: Minicursos", "Saiba mais: Mesas temáticas"… Só ajuda se ela usar a lista de links ou o Tab; por cabeçalhos ou pelo menu CHAMADAS a diferença some.',
         riscos='A página inicial chama o modelo para o logo: com o modelo padrão (qwen3:8b, raciocínio original) o alt gravado pode ser o raciocínio do modelo (D-8). Registrar a estratégia usada.'),
    # ---------------- Complementar ----------------
    dict(id='SBC-1', conj='Complementar', grupo='G2', site='SBC — Associe-se', url='https://www.sbc.org.br/',
         titulo='Valor da anuidade de uma categoria (tabela sem cabeçalho)', tipo='positivo', pag=2, est_s=90, extra=0,
         instr_a='No site da SBC, descubra qual é o valor atual da anuidade para a categoria Associado ou Associada Estudante de Pós-graduação.',
         saida_a='R$ 142,00',
         crit_a='S: R$ 142,00. P: chega à tabela de anuidades, mas lê o valor de outra linha. F: valor desatualizado da notícia de 2017 (R$ 70, 75, 175 ou 195), não chega à tabela, desistência.',
         instr_b='No site da SBC, descubra qual é o valor atual da anuidade para a categoria Associado ou Associada Efetivo/Fundador.',
         saida_b='R$ 394,00',
         crit_b='S: R$ 394,00. P: chega à tabela, mas lê o valor de outra linha. F: valor desatualizado da notícia de 2017 (R$ 70, 75, 175 ou 195), não chega à tabela, desistência.',
         caminho='Lista de links → "Sua Associação" (ou "Saiba mais" de "Seja Associado(a)") → T até a tabela (10×3) → linha da categoria → coluna 3.',
         efeito='Baixo, POSITIVO. A 1ª linha da tabela vira th, e o NVDA passa a anunciar "Valor da Anuidade" ao mudar de coluna; "Saiba mais" ganha contexto. Do lado negativo: o skip link inserido aponta para o topo da página, o logo vira "Abrir link 1" e a página inicial chama o modelo para imagens.',
         riscos='A SBC pode publicar os valores de 2027 na campanha de renovação (dezembro, pelo histórico): rodar o conferir na véspera. Registrar se ela usou a busca do site.'),
    dict(id='SBC-2', conj='Complementar', grupo='G2', site='SBC — Secretarias Regionais', url='https://www.sbc.org.br/',
         titulo='Quem é o secretário regional de um estado (tabela sem cabeçalho)', tipo='positivo', pag=2, est_s=65, extra=0,
         instr_a='No site da SBC, descubra quem é o secretário ou a secretária regional da SBC em Pernambuco.',
         saida_a='Kiev Gama',
         crit_a='S: "Kiev Gama". P: lê outra coluna da linha certa, ou dá a secretária adjunta (Taciana Pontual Falcão, na página do estado). F: nome de outro estado, não chega à tabela, desistência.',
         instr_b='No site da SBC, descubra quem é o secretário ou a secretária regional da SBC em Roraima.',
         saida_b='Herbert Rocha',
         crit_b='S: "Herbert Rocha". P: lê outra coluna da linha certa. F: nome de outro estado, não chega à tabela, desistência.',
         caminho='Lista de links → "Sobre as Secretarias" → busca do NVDA pelo estado → Ctrl+Alt+seta para a direita (coluna 2).',
         efeito='Baixo, POSITIVO. A 1ª linha vira th e o NVDA anuncia "Secretário (a) Regional" ao entrar na coluna 2. Custos: skip link para o topo, rótulos genéricos no logo e na busca.',
         riscos='As páginas de cada estado podem divergir da tabela (na Paraíba divergem): a tabela é a referência. Rodar o conferir na véspera.'),
    dict(id='EDU-3', conj='Complementar', grupo='G1', site='EduComp 2026 — Comitê de programa', url='https://educomp.sbc.org.br/2026/educomp/equipes/comite-programa.html',
         titulo='Instituições dos coordenadores de uma trilha (tabela de dados sem th)', tipo='positivo', pag=1, est_s=35, extra=0,
         instr_a='No comitê de programa do EduComp 2026, descubra de quais instituições são os coordenadores da Trilha 2, Relatos de experiência.',
         saida_a='UEA e UFRR (Fernanda Pires — UEA; Luciano Ferreira — UFRR)',
         crit_a='S: as duas instituições. P: só uma, ou os nomes sem as instituições. F: instituições de outra trilha, desistência.',
         instr_b='No comitê de programa do EduComp 2026, descubra de quais instituições são os coordenadores da Trilha 4, Recursos e ambientes educacionais.',
         saida_b='UFMG e UFJF (Eduardo Figueiredo — UFMG; Ronney Castro — UFJF)',
         crit_b='S: as duas instituições. P: só uma, ou os nomes sem as instituições. F: instituições de outra trilha, desistência.',
         caminho='Lista de cabeçalhos → "Trilha 2 - Relatos de experiência" (B: "Trilha 4 - …") → T → ler as duas linhas.',
         efeito='Baixo, POSITIVO. A linha "Nome | Instituição" vira th e o NVDA anuncia "Instituição" ao mudar de coluna. Tabela pequena (3×2): lendo a linha inteira, a diferença some.',
         riscos='Tarefa curta: a operação do painel pesa proporcionalmente mais.'),
    dict(id='JIS-1', conj='Complementar', grupo='G1', site='JIS — volume 16 (2025)', url='https://journals-sol.sbc.org.br/index.php/jis',
         titulo='Páginas de um artigo no volume de 2025 do JIS', tipo='controle', pag=3, est_s=70, extra=0,
         instr_a='No Journal on Interactive Systems, descubra em quais páginas do volume de 2025 está publicado o artigo "Kahoot as a data collection tool to work with school age children".',
         saida_a='227–235',
         crit_a='S: 227 a 235. P: chega ao artigo, mas dá as páginas de um vizinho. F: não chega ao volume de 2025, desistência.',
         instr_b='No Journal on Interactive Systems, descubra em quais páginas do volume de 2025 está publicado o artigo "Evaluating a Serious Game for Math: Is it Useful and Fun?".',
         saida_b='199–212',
         crit_b='S: 199 a 212. P: chega ao artigo, mas dá as páginas de um vizinho. F: não chega ao volume de 2025, desistência.',
         caminho='"Regular Issues" (arquivo) → "Vol. 16 No. 1 (2025)" → busca do NVDA por "Kahoot" (B: "Math") → seta para baixo (autores, páginas).',
         efeito='Nenhum a baixo. A inicial do JIS teve 146 violações no axe, a pior entre as páginas iniciais do estrato D (a Lista A–Z da CAPES, página de tarefa, teve 172), mas para leitor de tela é bem estruturada: o número superestima a barreira. A extensão mexe em contraste, no logo e em métricas da barra lateral, fora do caminho; o logo sem alt vai ao modelo em toda página.',
         riscos='A busca do JIS leva à SOL (outro servidor). Página em inglês: o NVDA pode trocar de voz, igual nas duas condições.'),
    # ---------------- Reservas ----------------
    dict(id='SOL-4', conj='Reserva', grupo='—', site='SOL Livros — série de IHC', url='https://books-sol.sbc.org.br/index.php/sbc/catalog/series/ppcihc',
         titulo='ISBN de um livro da série de IHC (substitui SOL-1 se a SOL estiver fora do ar)', tipo='baixo', pag=2, est_s=60, extra=0,
         instr_a='Na série de livros "Perspectivas e Práticas Contemporâneas em IHC", da SOL, encontre o livro "Minicursos do IHC 2024" e diga o ISBN dele.',
         saida_a='978-85-7669-623-0',
         crit_a='S: o ISBN completo (hífens opcionais). P: página certa, com um dígito trocado, ou o DOI no lugar do ISBN. F: ISBN de outro livro (registrar se seguiu um link de capa), desistência.',
         instr_b='Na série de livros "Perspectivas e Práticas Contemporâneas em IHC", da SOL, encontre o livro "Minicursos do IHC 2023" e diga o ISBN dele.',
         saida_b='978-85-7669-604-9',
         crit_b='S: o ISBN completo. P: página certa, com um dígito trocado, ou o DOI. F: ISBN de outro livro, desistência.',
         caminho='Tecla 3 até o título do livro → página do livro → cabeçalho "ISBN-13 (15)".',
         efeito='Baixo, possivelmente negativo. As capas são links sem nome; a extensão as rotula com nome genérico (painel) ou com o cabeçalho do livro anterior (modo automático), o que pode levar ao livro errado pela lista de links.',
         riscos='Servidor books-sol, diferente do sol.sbc.org.br. Ler ISBN em voz alta é sujeito a erro de dígito. No modo automático, cada capa herda o título do livro anterior: pela lista de links, "Minicursos do IHC 2023" pode levar ao livro de 2024 (B) e "Minicursos do IHC 2024" a "Práticas de IHC em Sala de Aula" (A); classificar como falha induzida pela extensão.'),
    dict(id='CAPES-3', conj='Reserva', grupo='—', site='gov.br/capes — Valores de Bolsas', url='https://www.gov.br/capes/pt-br/acesso-a-informacao/acoes-e-programas/bolsas',
         titulo='Valor de uma bolsa no exterior (tabela)', tipo='baixo', pag=2, est_s=110, extra=0,
         instr_a='No site da CAPES, descubra o valor mensal da bolsa de Professor Visitante Sênior no exterior, em dólar canadense.',
         saida_a='3.060,00 dólares canadenses',
         crit_a='S: 3.060 (também vale o mesmo valor lido na seção histórica "Professor/Pesquisador Visitante no Exterior", linha Mensalidade). P: linha certa e coluna vizinha (3.420,00 ou 1.900,00), ou moeda errada. F: valor em reais (bolsa no país), desistência.',
         instr_b='No site da CAPES, descubra o valor mensal da bolsa de Professor Visitante Júnior Pós-Doutorado no exterior, em dólar canadense.',
         saida_b='2.660,00 dólares canadenses',
         crit_b='S: 2.660 (também vale o valor lido em "Estágio Pós-Doutoral", linha Mensalidade). P: linha certa e coluna vizinha (3.000,00 ou 1.700,00), ou moeda errada. F: valor em reais, desistência.',
         caminho='"Prestação de Contas" → "Valores de Bolsas" → 2ª tabela ("Valores das bolsas no exterior") → linha → coluna "Dólar Canadense".',
         efeito='Baixo. As tabelas já têm th; o que atrapalha (títulos de seção que não são cabeçalhos, sumário com âncoras quebradas) não é tratado. 5 links vazios viram "Abrir link N".',
         riscos='Página com 44 tabelas; gov.br devolveu 429 após poucas requisições na exploração.'),
    dict(id='IHC-3', conj='Reserva', grupo='—', site='IHC 2025 — página inicial', url='https://ihc.sbc.org.br/2025/',
         titulo='Dia e hora de início de uma sessão especial (substitui CSBC-2)', tipo='baixo', pag=2, est_s=65, extra=0,
         instr_a='No site do IHC 2025, descubra em que dia e a que horas começou a sessão de apresentações do Prêmio Junia Coutinho Anacleto.',
         saida_a='Terça-feira, 09/09/2025, às 8h',
         crit_a='S: dia (terça ou 09/09) + início às 8h. P: só o dia ou só a hora. F: outro dia ou hora, desistência.',
         instr_b='No site do IHC 2025, descubra em que dia e a que horas começou a Sessão 1 de Pôsteres e Demonstrações.',
         saida_b='Segunda-feira, 08/09/2025, às 15h30',
         crit_b='S: dia (segunda ou 08/09) + início às 15h30. P: só o dia ou só a hora. F: dados da Sessão 2 (09/09), "data a ser definida", desistência.',
         caminho='Cabeçalho do cartão → "Saiba mais" → cabeçalho "Programação" → h4 com dia e hora.',
         efeito='Baixo. Os 3 "Saiba mais" iguais ganham contexto (ajuda pela lista de links); as sobreposições dos cartões viram "Abrir link N".',
         riscos='Página da amostra do Estudo 1 (axe 19 → 2). Rota alternativa pela Programação dá as mesmas respostas: registrar a rota.'),
    dict(id='IHC-4', conj='Reserva', grupo='—', site='IHC 2026 — Local', url='https://ihc.sbc.org.br/2026/',
         titulo='Distância de um hotel recomendado até a UFC', tipo='baixo', pag=2, est_s=70, extra=0,
         instr_a='No site do IHC 2026, descubra a que distância da UFC fica o Hotel Pedra dos Ventos, segundo a lista de hotéis do evento.',
         saida_a='23 km',
         crit_a='S: 23 km. P: acha o hotel, mas não a distância. F: distância de outro hotel, desistência.',
         instr_b='No site do IHC 2026, descubra a que distância da UFC fica o Cedro Palace Hotel, segundo a lista de hotéis do evento.',
         saida_b='5,2 km',
         crit_b='S: 5,2 km. P: acha o hotel, mas não a distância, ou arredonda para 5 km. F: distância de outro hotel, desistência.',
         caminho='"Local" → busca do NVDA pelo nome do hotel (ou tecla 4) → lista do cartão, 4º item.',
         efeito='Baixo, provavelmente negativo: ícones decorativos viram "gráfico" no cartão; links de mapa só com ícone viram "Abrir link N".',
         riscos='A página Local pode ser editada ou sair do menu depois do evento.'),
    dict(id='IHC-5', conj='Reserva', grupo='—', site='IHC 2025 — Inscrições', url='https://ihc.sbc.org.br/2025/',
         titulo='Valor de uma categoria de inscrição (somente IHC, 1º prazo)', tipo='baixo', pag=2, est_s=100, extra=0,
         instr_a='No site do IHC 2025, descubra o valor da inscrição somente no IHC, sem o INTERACT, para a categoria Profissional associado à SBC com anuidade vigente, no prazo até 7 de julho.',
         saida_a='R$ 450,00',
         crit_a='S: R$ 450,00. P: categoria certa, outro prazo. F: valor do Combo IHC + INTERACT ou de outra categoria, desistência.',
         instr_b='No site do IHC 2025, descubra o valor da inscrição somente no IHC, sem o INTERACT, para a categoria Estudante de pós-graduação associado à SBC com anuidade vigente, no prazo até 7 de julho.',
         saida_b='R$ 290,00',
         crit_b='S: R$ 290,00. P: categoria certa, outro prazo (R$ 320,00, 390,00 ou 440,00). F: Combo (R$ 700,00), outra categoria, desistência.',
         caminho='"Inscrições" → cabeçalho "Somente Inscrição no IHC" → h4 da categoria → h5 "… com anuidade vigente" → "Até 07/07".',
         efeito='Baixo: a dificuldade é de conteúdo (Combo e Somente IHC repetem as categorias). A extensão acrescenta "gráfico" perto dos cartões.',
         riscos='Instrução com 4 restrições: permitir repetir e contar as repetições.'),
    dict(id='CSBC-1', conj='Reserva', grupo='—', site='CSBC 2026 — Programação Geral', url='https://csbc.sbc.org.br/2026/',
         titulo='Horário de uma atividade na programação (controle)', tipo='controle', pag=2, est_s=75, extra=0,
         instr_a='No site do CSBC 2026, descubra a que horas começou a Cerimônia de Abertura, que aconteceu na segunda-feira, 20 de julho.',
         saida_a='19h (19:00–20:00, Sala Hortênsia Central)',
         crit_a='S: 19h. P: lê o horário vizinho (18:30 ou 20:30) ou confunde com o Coquetel de Abertura. F: outro horário sem localizar o item, desistência.',
         instr_b='No site do CSBC 2026, descubra a que horas começou o Jantar do Evento, que aconteceu na quarta-feira, 22 de julho.',
         saida_b='20h',
         crit_b='S: 20h (Programação Geral ou página do Jantar). P: lê o horário vizinho, ou 19h da página do WIT (fonte divergente do próprio site). F: outro horário sem localizar o item, desistência.',
         caminho='"Programação" → aba do dia → busca do NVDA por "Cerimônia" (B: "Jantar") → seta para cima até o horário.',
         efeito='Nenhum (controle): abas e itens já têm nome.',
         riscos='Itens da programação são "botões" que não fazem nada.'),
    dict(id='JIS-2', conj='Reserva', grupo='—', site='JIS — diretrizes para autores', url='https://journals-sol.sbc.org.br/index.php/jis',
         titulo='Uma regra das diretrizes do JIS (controle; substitui SOL-2 se a SOL estiver fora do ar)', tipo='controle', pag=2, est_s=60, extra=0,
         instr_a='Nas diretrizes para autores do Journal on Interactive Systems, descubra quantas contribuições originais, no máximo, a carta de apresentação deve listar.',
         saida_a='Até cinco',
         crit_a='S: "até cinco", depois de a gravação mostrar o NVDA lendo o trecho. P: chega às diretrizes e dá outro número. F: não encontra as diretrizes, desistência.',
         instr_b='Nas diretrizes para autores do Journal on Interactive Systems, descubra por quantos especialistas independentes os manuscritos submetidos costumam ser avaliados.',
         saida_b='Três',
         crit_b='S: "três", depois de a gravação mostrar o NVDA lendo o trecho. P: chega às diretrizes e dá outro número. F: não encontra as diretrizes, desistência.',
         caminho='Menu "Submissions" → "Authors Guidelines" → cabeçalho "Cover Letter" (B: "Review Process").',
         efeito='Nenhum (controle): cabeçalhos e listas corretos; menu já acessível.',
         riscos='A coautora pode conhecer as diretrizes: exigir que o trecho seja lido na gravação.'),
]

for t in T:
    t['est_c'] = t['est_s'] + PAINEL_S_POR_PAGINA * t['pag'] + t['extra']
TAREFA = {t['id']: t for t in T}

# Ordem da sessão: grupo G1 começa sem extensão, G2 começa com extensão; a 1ª exposição é sempre a variante A.
NUCLEO_G1 = ['SOL-1', 'CAPES-1', 'IHC-2', 'CSBC-2']
NUCLEO_G2 = ['SOL-2', 'CAPES-2', 'IHC-1', 'EDU-1']
COMPL_G1 = ['JIS-1', 'EDU-3']
COMPL_G2 = ['SBC-1', 'SBC-2']
for _lista in (NUCLEO_G1, NUCLEO_G2, COMPL_G1, COMPL_G2):
    for _i, _tid in enumerate(_lista):
        if _tid in TROCAS:
            _res = TAREFA[TROCAS[_tid]]
            assert _res['conj'] == 'Reserva', f'{TROCAS[_tid]} não é reserva'
            _res.update(conj=f"{TAREFA[_tid]['conj']} (substitui {_tid})", grupo=TAREFA[_tid]['grupo'])
            _lista[_i] = _res['id']
BLOCOS = [
    ('Sessão 1 · Bloco 1a', 'S', 'A', NUCLEO_G1),
    ('Sessão 1 · Bloco 1b', 'C', 'A', NUCLEO_G2),
    ('Sessão 1 · Bloco 2a', 'C', 'B', NUCLEO_G1),
    ('Sessão 1 · Bloco 2b', 'S', 'B', NUCLEO_G2),
    ('Sessão 2 · Bloco 3a', 'S', 'A', COMPL_G1),
    ('Sessão 2 · Bloco 3b', 'C', 'A', COMPL_G2),
    ('Sessão 2 · Bloco 4a', 'C', 'B', COMPL_G1),
    ('Sessão 2 · Bloco 4b', 'S', 'B', COMPL_G2),
]
P0 = dict(id='P0', site='W3C BAD (controle do Estudo 1)', url='https://www.w3.org/WAI/demos/bad/before/home.html',
          instr='Com o Accessify, analise esta página, aplique as otimizações e volte para a página. Avise quando estiver de volta.',
          saida='Painel anuncia que as otimizações foram aplicadas; o foco volta à página',
          crit='S: aciona Alt+Shift+A, chega ao painel, ouve o resultado da análise, ativa "Aplicar Otimizações", ouve a confirmação e volta à página, sem ajuda. P: conclui com uma dica. F: não conclui. (Treino: a dica é permitida, mas registrada.)',
          est=90)

CABECALHOS_BLOCO = [
    'Sessão 1 · Preparação e regras (15 min); familiarização livre no perfil S (2 min) · Bloco 1a — grupo G1, SEM extensão (perfil S), variante A',
    'Trocar para o perfil C · P0, treino do painel · Bloco 1b — grupo G2, COM extensão, variante A',
    'Pausa (10 min) · Bloco 2a — grupo G1, COM extensão (perfil C), variante B',
    'Trocar para o perfil S · Bloco 2b — grupo G2, SEM extensão, variante B · depois, entrevista final (10 min)',
    'Sessão 2 (complementar) · Preparação (10 min) · Bloco 3a — grupo G1, SEM extensão (perfil S), variante A',
    'Trocar para o perfil C · Bloco 3b — grupo G2, COM extensão, variante A',
    'Bloco 4a — grupo G1, COM extensão (perfil C), variante B',
    'Trocar para o perfil S · Bloco 4b — grupo G2, SEM extensão, variante B · entrevista curta (5 min)',
]

# ------------------------------------------------------------------ xlsx
ARIAL = 'Arial'
F_BASE = Font(name=ARIAL, size=10)
F_BOLD = Font(name=ARIAL, size=10, bold=True)
F_HEAD = Font(name=ARIAL, size=10, bold=True, color='FFFFFF')
F_TITLE = Font(name=ARIAL, size=13, bold=True)
F_EX = Font(name=ARIAL, size=10, italic=True, color='7F7F7F')
F_ALERTA = Font(name=ARIAL, size=10, bold=True, color='C00000')
FILL_HEAD = PatternFill('solid', fgColor='305496')
FILL_INPUT = PatternFill('solid', fgColor='FFF2CC')
FILL_CALC = PatternFill('solid', fgColor='EDEDED')
FILL_BLOCK = PatternFill('solid', fgColor='D9E1F2')
FILL_S = PatternFill('solid', fgColor='FCE4D6')
FILL_C = PatternFill('solid', fgColor='E2EFDA')
THIN = Side(style='thin', color='BFBFBF')
BORDA = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='top', wrap_text=True)


def cabecalho(ws, linha, nomes_larguras):
    for j, (h, w) in enumerate(nomes_larguras, 1):
        c = ws.cell(row=linha, column=j, value=h)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = CENTER
        c.border = BORDA
        if w:
            ws.column_dimensions[get_column_letter(j)].width = w


wb = Workbook()

# --- Roteiro
ws = wb.active
ws.title = 'Roteiro'
COLS = [
    ('Ordem', 7), ('Bloco', 18), ('Condição', 10), ('Tarefa', 9), ('Variante', 8), ('Site', 22), ('URL inicial', 30),
    ('Ação (ler em voz alta)', 48), ('Saída esperada', 30), ('Critério de sucesso', 44), ('Estimativa (s)', 10),
    ('Resposta dada', 26), ('Resultado (S/P/F)', 10), ('t0 no vídeo (mm:ss)', 11), ('t1 no vídeo (mm:ss)', 11),
    ('t0 (s)', 8), ('t1 (s)', 8), ('Tempo total (s)', 10), ('Tempo no painel (s)', 10), ('Tempo líquido (s)', 10),
    ('Acima do teto (300 s líquidos)?', 11), ('Encerrada pelo(a) facilitador(a) (s/n)', 12), ('Repetições da instrução', 10),
    ('Lembretes para acionar o Accessify', 12), ('Ajuda (sim → F)', 9), ('SEQ (1–7)', 8), ('Estratégia', 18),
    ('Barreiras e falas do NVDA', 40), ('Observações', 30), ('Conferência', 26), ('Chave', 12),
]
IDX = {nome: i + 1 for i, (nome, _) in enumerate(COLS)}
L = {nome: get_column_letter(i) for nome, i in IDX.items()}
INPUTS = ['Resposta dada', 'Resultado (S/P/F)', 't0 no vídeo (mm:ss)', 't1 no vídeo (mm:ss)', 'Tempo no painel (s)',
          'Encerrada pelo(a) facilitador(a) (s/n)', 'Repetições da instrução', 'Lembretes para acionar o Accessify',
          'Ajuda (sim → F)', 'SEQ (1–7)', 'Estratégia', 'Barreiras e falas do NVDA', 'Observações']
CALCS = ['t0 (s)', 't1 (s)', 'Tempo total (s)', 'Tempo líquido (s)', 'Acima do teto (300 s líquidos)?', 'Conferência', 'Chave']
CENTRALIZADAS = ['Ordem', 'Condição', 'Tarefa', 'Variante', 'Estimativa (s)', 'Resultado (S/P/F)', 't0 no vídeo (mm:ss)',
                 't1 no vídeo (mm:ss)', 'Tempo no painel (s)', 'Encerrada pelo(a) facilitador(a) (s/n)',
                 'Repetições da instrução', 'Lembretes para acionar o Accessify', 'Ajuda (sim → F)', 'SEQ (1–7)']

ws['A1'] = 'Estudo 2 — Roteiro e registro das tarefas (uma linha por execução)'
ws['A1'].font = F_TITLE
ws['A2'] = ('Preencher só as células amarelas, a partir do vídeo e do log do NVDA; as cinzas são calculadas e a coluna '
            'Conferência aponta inconsistências. Tempos do vídeo como mm:ss ou h:mm:ss. Condição: S = sem extensão, '
            'C = com Accessify v1.1.0. Definições na aba Instruções. A linha 4 é um exemplo e não entra no Resumo.')
ws['A2'].font = F_BASE
ws['A2'].alignment = WRAP
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=12)
ws.row_dimensions[2].height = 42
HR = 3
cabecalho(ws, HR, COLS)
ws.row_dimensions[HR].height = 54


def ref(nome, r):
    return f'{L[nome]}{r}'


def seg(c):
    """Texto mm:ss ou h:mm:ss → segundos."""
    return (f'IF({c}="","",IF(LEN({c})-LEN(SUBSTITUTE({c},":",""))=2,'
            f'LEFT({c},FIND(":",{c})-1)*3600+MID({c},FIND(":",{c})+1,2)*60+RIGHT({c},2),'
            f'LEFT({c},FIND(":",{c})-1)*60+RIGHT({c},2)))')


def formulas(r):
    tot, pai, liq = ref('Tempo total (s)', r), ref('Tempo no painel (s)', r), ref('Tempo líquido (s)', r)
    teto, res = ref('Acima do teto (300 s líquidos)?', r), ref('Resultado (S/P/F)', r)
    enc, aju, cond = ref('Encerrada pelo(a) facilitador(a) (s/n)', r), ref('Ajuda (sim → F)', r), ref('Condição', r)
    t0, t1 = ref('t0 (s)', r), ref('t1 (s)', r)
    return {
        't0 (s)': '=' + seg(ref('t0 no vídeo (mm:ss)', r)),
        't1 (s)': '=' + seg(ref('t1 no vídeo (mm:ss)', r)),
        'Tempo total (s)': f'=IF(OR({t0}="",{t1}=""),"",{t1}-{t0})',
        'Tempo líquido (s)': f'=IF({tot}="","",IF(AND({pai}<>"",NOT(ISNUMBER({pai}))),"painel inválido",{tot}-N({pai})))',
        'Acima do teto (300 s líquidos)?': f'=IF(NOT(ISNUMBER({liq})),"",IF({liq}>300,"sim","não"))',
        'Conferência': (f'=TRIM(IF(AND(ISNUMBER({tot}),{tot}<=0),"t1 não é depois de t0; ","")'
                        f'&IF(AND({cond}="C",ISNUMBER({tot}),{pai}=""),"falta o tempo no painel; ","")'
                        f'&IF(AND({teto}="sim",{res}<>"F"),"acima do teto: deve ser F; ","")'
                        f'&IF(AND({enc}="s",{res}<>"F"),"encerrada: deve ser F; ","")'
                        f'&IF(AND({aju}="sim",{res}<>"F"),"houve ajuda: deve ser F; ","")'
                        f'&IF(AND({res}<>"",ISNUMBER({tot})=FALSE),"falta t0 ou t1; ",""))'),
        'Chave': f'={ref("Tarefa", r)}&"|"&{ref("Condição", r)}',
    }


def escreve(r, valores, fill_cond=None, fonte=F_BASE):
    for nome, v in valores.items():
        ws.cell(row=r, column=IDX[nome], value=v)
    for nome, v in formulas(r).items():
        ws.cell(row=r, column=IDX[nome], value=v)
    for nome in IDX:
        c = ws.cell(row=r, column=IDX[nome])
        c.border = BORDA
        c.font = F_ALERTA if nome == 'Conferência' else fonte
        c.alignment = CENTER if nome in CENTRALIZADAS or nome in CALCS[:5] else WRAP
        if nome in INPUTS:
            c.fill = FILL_INPUT
        elif nome in CALCS:
            c.fill = FILL_CALC
        elif nome == 'Condição' and fill_cond is not None:
            c.fill = fill_cond
    for nome in ('t0 no vídeo (mm:ss)', 't1 no vídeo (mm:ss)'):
        ws.cell(row=r, column=IDX[nome]).number_format = '@'


r = HR + 1
ex = TAREFA['IHC-2']
escreve(r, {
    'Ordem': 'ex.', 'Bloco': 'EXEMPLO (não conta)', 'Condição': '—', 'Tarefa': 'IHC-2', 'Variante': 'A', 'Site': ex['site'],
    'URL inicial': ex['url'], 'Ação (ler em voz alta)': ex['instr_a'], 'Saída esperada': ex['saida_a'],
    'Critério de sucesso': ex['crit_a'], 'Estimativa (s)': ex['est_s'],
    'Resposta dada': 'Sessão técnica 2, Diversidade e Justiça Social', 'Resultado (S/P/F)': 'S',
    't0 no vídeo (mm:ss)': '12:05', 't1 no vídeo (mm:ss)': '13:21', 'Tempo no painel (s)': 0,
    'Encerrada pelo(a) facilitador(a) (s/n)': 'n', 'Repetições da instrução': 0, 'Lembretes para acionar o Accessify': 0,
    'Ajuda (sim → F)': 'não', 'SEQ (1–7)': 6, 'Estratégia': 'várias',
    'Barreiras e falas do NVDA': '12:31 buscou "10h30" com a aba de segunda aberta e caiu na linha errada; voltou e ativou a aba "20/10 - Terça".',
    'Observações': 'Usou Ctrl+Alt+setas na tabela.',
}, fonte=F_EX)
ws.cell(row=r, column=IDX['t0 no vídeo (mm:ss)']).comment = Comment('t0: o(a) facilitador(a) diz "pode começar".', 'Estudo 2')
ws.cell(row=r, column=IDX['t1 no vídeo (mm:ss)']).comment = Comment('t1: a avaliadora diz a resposta final (vale a última) ou "desisto".', 'Estudo 2')
r += 1
primeira = r


def linha_bloco(texto):
    global r
    c = ws.cell(row=r, column=1, value=texto)
    c.font = F_BOLD
    for col in range(1, len(COLS) + 1):
        ws.cell(row=r, column=col).fill = FILL_BLOCK
    r += 1


ordem = 1
linha_exposicao_a = {}
for i, (bloco, cond, var, ids) in enumerate(BLOCOS):
    linha_bloco(CABECALHOS_BLOCO[i])
    if i == 1:
        escreve(r, {'Ordem': 0, 'Bloco': 'Sessão 1 · P0 (treino)', 'Condição': 'Painel', 'Tarefa': 'P0', 'Variante': '—',
                    'Site': P0['site'], 'URL inicial': P0['url'], 'Ação (ler em voz alta)': P0['instr'],
                    'Saída esperada': P0['saida'], 'Critério de sucesso': P0['crit'], 'Estimativa (s)': P0['est']})
        r += 1
    for tid in ids:
        t = TAREFA[tid]
        v = var.lower()
        if var == 'A':
            linha_exposicao_a[tid] = r
        escreve(r, {
            'Ordem': ordem, 'Bloco': bloco, 'Condição': cond, 'Tarefa': tid, 'Variante': var, 'Site': t['site'],
            'URL inicial': t['url'], 'Ação (ler em voz alta)': t['instr_' + v], 'Saída esperada': t['saida_' + v],
            'Critério de sucesso': t['crit_' + v], 'Estimativa (s)': t['est_s'] if cond == 'S' else t['est_c'],
        }, fill_cond=FILL_S if cond == 'S' else FILL_C)
        ordem += 1
        r += 1
ultima = r - 1
ws.freeze_panes = ws.cell(row=HR + 1, column=IDX['Site'])
ws.auto_filter.ref = f'A{HR}:{get_column_letter(len(COLS))}{ultima}'


def valida(dv, nome, msg=None):
    dv.showErrorMessage = msg is not None
    if msg:
        dv.errorTitle = 'Valor fora do formato'
        dv.error = msg
    ws.add_data_validation(dv)
    dv.add(f'{L[nome]}{primeira}:{L[nome]}{ultima}')


valida(DataValidation(type='list', formula1='"S,P,F"', allow_blank=True), 'Resultado (S/P/F)', 'Use S (sucesso), P (parcial) ou F (falha).')
valida(DataValidation(type='whole', operator='between', formula1='1', formula2='7', allow_blank=True), 'SEQ (1–7)', 'SEQ vai de 1 (muito difícil) a 7 (muito fácil).')
valida(DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True), 'Tempo no painel (s)', 'Tempo no painel em segundos (número ≥ 0).')
valida(DataValidation(type='whole', operator='greaterThanOrEqual', formula1='0', allow_blank=True), 'Repetições da instrução', 'Número inteiro ≥ 0.')
valida(DataValidation(type='whole', operator='greaterThanOrEqual', formula1='0', allow_blank=True), 'Lembretes para acionar o Accessify', 'Número inteiro ≥ 0.')
valida(DataValidation(type='list', formula1='"s,n"', allow_blank=True), 'Encerrada pelo(a) facilitador(a) (s/n)', 'Use s ou n.')
valida(DataValidation(type='list', formula1='"sim,não"', allow_blank=True), 'Ajuda (sim → F)', 'Use sim ou não.')
valida(DataValidation(type='list', formula1='"cabeçalhos,lista de links,busca do NVDA,leitura linear,tabela,menu do site,busca do site,várias"', allow_blank=True), 'Estratégia')
for nome in ('t0 no vídeo (mm:ss)', 't1 no vídeo (mm:ss)'):
    c = f'{L[nome]}{primeira}'
    n_dois_pontos = f'LEN({c})-LEN(SUBSTITUTE({c},":",""))'
    formula = (f'IF({c}="",TRUE,IFERROR(AND({c}=TRIM({c}),ISERROR(FIND(".",{c})),MID({c},LEN({c})-2,1)=":",'
               f'ISNUMBER(-RIGHT({c},2)),ISNUMBER(-LEFT({c},FIND(":",{c})-1)),'
               f'IF({n_dois_pontos}=1,TRUE,IF({n_dois_pontos}=2,MID({c},LEN({c})-5,1)=":",FALSE))),FALSE))')
    valida(DataValidation(type='custom', formula1=formula, allow_blank=True), nome,
           'Use mm:ss ou h:mm:ss, sem frações de segundo nem quadros (ex.: 12:05 ou 1:02:03).')
ws.page_setup.orientation = 'landscape'
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0

# --- Tarefas (banco)
wt = wb.create_sheet('Tarefas')
CABT = [('ID', 9), ('Conjunto', 14), ('Grupo', 7), ('Site', 22), ('URL inicial', 30), ('Objetivo', 30), ('Efeito previsto', 11),
        ('Por que (previsão)', 50), ('Páginas no caminho', 9), ('Estimativa S (s)', 9), ('Estimativa C (s)', 9),
        ('Ação A', 45), ('Saída esperada A', 26), ('Critério A', 40), ('Ação B', 45), ('Saída esperada B', 26),
        ('Critério B', 40), ('Caminho esperado', 40), ('Riscos e cuidados', 45)]
cabecalho(wt, 1, CABT)
for i, t in enumerate(T):
    vals = [t['id'], t['conj'], t['grupo'], t['site'], t['url'], t['titulo'], t['tipo'], t['efeito'], t['pag'],
            t['est_s'], t['est_c'], t['instr_a'], t['saida_a'], t['crit_a'], t['instr_b'], t['saida_b'], t['crit_b'],
            t['caminho'], t['riscos']]
    for j, v in enumerate(vals, 1):
        c = wt.cell(row=2 + i, column=j, value=v.replace('`', '') if isinstance(v, str) else v)
        c.font = F_BASE
        c.alignment = WRAP
        c.border = BORDA
ULT_T = len(T) + 1
wt.cell(row=ULT_T + 2, column=1, value=(
    f'Estimativa C = estimativa S + {PAINEL_S_POR_PAGINA} s por página no caminho (operação do painel lateral: Alt+Shift+A, '
    'ouvir o resultado, "Aplicar Otimizações", voltar à página) + contorno previsto (CSBC-2: +50 s). Estimativas S: tempo de '
    'uma usuária experiente de NVDA, revisto pelo verificador de cada tarefa. No modo automático, somar ≈ 3 s por página.')).font = F_EX
wt.freeze_panes = 'B2'

# --- Resumo
rs = wb.create_sheet('Resumo')
rs['A1'] = 'Resumo (calculado a partir do Roteiro)'
rs['A1'].font = F_TITLE
rs['A2'] = ('Uma avaliadora: relatar os valores por tarefa e as contagens; médias são só descritivas. Só entram linhas com '
            'resultado S, P ou F. Tempos em segundos.')
rs['A2'].font = F_BASE
RG = lambda nome: f'Roteiro!${L[nome]}${primeira}:${L[nome]}${ultima}'
TG = lambda col: f'Tarefas!${col}$2:${col}${ULT_T}'
cabecalho(rs, 4, [('Condição', 12), ('Execuções com resultado', 14), ('Sucesso (S)', 8), ('Parcial (P)', 11),
                  ('Falha (F)', 11), ('Taxa de sucesso', 11), ('Tempo total médio', 11), ('Tempo líquido médio', 11),
                  ('SEQ médio', 11), ('Acima do teto', 12), (None, None), (None, None), (None, None), (None, None)])
for i, cond in enumerate(['S', 'C']):
    rr = 5 + i
    cnd, res = RG('Condição'), RG('Resultado (S/P/F)')

    def media(col):
        # média só das linhas com resultado S, P ou F (sem curingas, que dependem da configuração do Calc)
        x = RG(col)
        soma = '+'.join(f'SUMIFS({x},{cnd},$A{rr},{res},"{k}")' for k in 'SPF')
        n = '+'.join(f'COUNTIFS({cnd},$A{rr},{res},"{k}",{x},">-1E+300")' for k in 'SPF')
        return f'=IFERROR(({soma})/({n}),"")'
    vals = [
        cond,
        f'=COUNTIFS({cnd},$A{rr},{res},"S")+COUNTIFS({cnd},$A{rr},{res},"P")+COUNTIFS({cnd},$A{rr},{res},"F")',
        f'=COUNTIFS({cnd},$A{rr},{res},"S")',
        f'=COUNTIFS({cnd},$A{rr},{res},"P")',
        f'=COUNTIFS({cnd},$A{rr},{res},"F")',
        f'=IF(B{rr}=0,"",C{rr}/B{rr})',
        media('Tempo total (s)'),
        media('Tempo líquido (s)'),
        media('SEQ (1–7)'),
        f'=COUNTIFS({cnd},$A{rr},{RG("Acima do teto (300 s líquidos)?")},"sim")',
    ]
    for j, v in enumerate(vals, 1):
        c = rs.cell(row=rr, column=j, value=v)
        c.font = F_BASE
        c.border = BORDA
        c.alignment = CENTER
        if j > 1:
            c.fill = FILL_CALC
    rs.cell(row=rr, column=6).number_format = '0%'
    for j in (7, 8, 9):
        rs.cell(row=rr, column=j).number_format = '0.0'

rs['A9'] = 'Por tarefa (variantes A e B nas duas condições; a tarefa vem do Roteiro, então uma troca feita lá aparece aqui)'
rs['A9'].font = F_BOLD
cabecalho(rs, 10, [(h, None) for h in ['Tarefa', 'Conjunto', 'Grupo', 'Efeito previsto', 'Resultado S', 'Resultado C',
                                        'Tempo total S', 'Tempo total C', 'Diferença C − S', 'Tempo líquido S',
                                        'Tempo líquido C', 'Diferença líquida C − S', 'SEQ S', 'SEQ C']])


def busca(col, rr, cond):
    x = f'INDEX({RG(col)},MATCH($A{rr}&"|{cond}",{RG("Chave")},0))'
    return f'=IFERROR(IF({x}="","",{x}),"")'


def da_tarefa(col_tarefas, rr):
    return f'=IFERROR(INDEX({TG(col_tarefas)},MATCH($A{rr},{TG("A")},0)),"")'


ids_resumo = NUCLEO_G1 + NUCLEO_G2 + COMPL_G1 + COMPL_G2
linha_resumo = {}
for i, tid in enumerate(ids_resumo):
    rr = 11 + i
    linha_resumo[tid] = rr
    vals = [f'=Roteiro!{L["Tarefa"]}{linha_exposicao_a[tid]}', da_tarefa('B', rr), da_tarefa('C', rr), da_tarefa('G', rr),
            busca('Resultado (S/P/F)', rr, 'S'), busca('Resultado (S/P/F)', rr, 'C'),
            busca('Tempo total (s)', rr, 'S'), busca('Tempo total (s)', rr, 'C'),
            f'=IF(AND(ISNUMBER(G{rr}),ISNUMBER(H{rr})),H{rr}-G{rr},"")',
            busca('Tempo líquido (s)', rr, 'S'), busca('Tempo líquido (s)', rr, 'C'),
            f'=IF(AND(ISNUMBER(J{rr}),ISNUMBER(K{rr})),K{rr}-J{rr},"")',
            busca('SEQ (1–7)', rr, 'S'), busca('SEQ (1–7)', rr, 'C')]
    for j, v in enumerate(vals, 1):
        c = rs.cell(row=rr, column=j, value=v)
        c.font = F_BASE
        c.border = BORDA
        c.alignment = CENTER
        if j > 1:
            c.fill = FILL_CALC
ra = 11 + len(ids_resumo) + 1
ctrl_g1, ctrl_g2 = linha_resumo[NUCLEO_G1[2]], linha_resumo[NUCLEO_G2[0]]
rs.cell(row=ra, column=1, value='Aprendizagem estimada pelos controles do núcleo (s, tempo líquido)').font = F_BOLD
c = rs.cell(row=ra, column=12, value=f'=IF(AND(ISNUMBER(L{ctrl_g1}),ISNUMBER(L{ctrl_g2})),(L{ctrl_g2}-L{ctrl_g1})/2,"")')
c.font = F_BASE
c.fill = FILL_CALC
c.border = BORDA
c.number_format = '0.0'
rs.cell(row=ra + 1, column=1, value=(
    'O controle do grupo G1 (S primeiro) tem C − S ≈ −aprendizagem; o do G2 (C primeiro), C − S ≈ +aprendizagem. '
    'Estimativa grosseira: duas tarefas diferentes, uma execução cada.')).font = F_EX
for j, w in enumerate([12, 16, 8, 14, 11, 11, 11, 11, 12, 11, 11, 12, 8, 8], 1):
    rs.column_dimensions[get_column_letter(j)].width = w
rs.row_dimensions[4].height = 30
rs.row_dimensions[10].height = 30

# --- Sessão
wsx = wb.create_sheet('Sessão')
wsx['A1'] = 'Dados da sessão (preencher antes de começar)'
wsx['A1'].font = F_TITLE
campos = [
    ('Data e local', ''), ('Avaliadora', ''), ('Facilitador(a)', ''), ('Codificador(a) 1 / 2', ''),
    ('Versão do NVDA', ''), ('NVDA: voz e velocidade', ''), ('NVDA: anunciar cabeçalhos de linha/coluna de tabela', ''),
    ('NVDA: ler automaticamente a página ao carregar', ''),
    ('NVDA: nível de log (Entrada/saída) e arquivo do log', ''), ('Versão do Chrome', ''),
    ('Perfil S (sem extensão) / Perfil C (com extensão)', ''),
    ('Build da extensão', 'v1.1.0-jis_submission (informar o commit)'),
    ('Provedor e modelo', 'Ollama local, qwen3:8b, raciocínio original (padrão da v1.1.0)'),
    ('Máquina do Ollama e latência medida no piloto', ''),
    ('Modo da condição C', 'painel lateral em cada página (recomendado) / modo automático'),
    ('Sequência de teclas do painel fixada no piloto', ''),
    ('Janela, zoom e largura útil da página com o painel aberto', 'maximizada (≥ 1280 px), zoom 100%'),
    ('Fontes de áudio gravadas (sistema/NVDA e microfone)', ''),
    ('Conferência das respostas (eval/tools/estudo2_conferir.py): data, hora e resultado', ''),
    ('Arquivo da gravação (OBS)', ''), ('Tarefas trocadas por reserva (e motivo)', ''),
]
for i, (k, v) in enumerate(campos):
    a = wsx.cell(row=3 + i, column=1, value=k)
    a.font = F_BOLD
    a.alignment = WRAP
    b = wsx.cell(row=3 + i, column=2, value=v or None)
    b.font = F_BASE
    b.fill = FILL_INPUT
    b.border = BORDA
    b.alignment = WRAP
base = 3 + len(campos) + 1
wsx.cell(row=base, column=1, value='Familiaridade prévia da avaliadora com cada site').font = F_BOLD
sites = ['SOL — anais e busca (sol.sbc.org.br)', 'SOL Livros (books-sol.sbc.org.br)', 'Portal de Periódicos CAPES',
         'Site da CAPES (gov.br/capes)', 'Site do IHC (2025/2026)', 'Site do CSBC 2026', 'Site do EduComp 2026',
         'Site da SBC (sbc.org.br)', 'Journal on Interactive Systems (JIS)']
dv_fam = DataValidation(type='list', formula1='"nunca usei,já usei,uso com frequência"', allow_blank=True)
dv_fam.showErrorMessage = True
dv_fam.error = 'Use: nunca usei, já usei ou uso com frequência.'
wsx.add_data_validation(dv_fam)
for i, s in enumerate(sites):
    wsx.cell(row=base + 1 + i, column=1, value=s).font = F_BASE
    b = wsx.cell(row=base + 1 + i, column=2)
    b.fill = FILL_INPUT
    b.border = BORDA
    dv_fam.add(b.coordinate)
wsx.column_dimensions['A'].width = 52
wsx.column_dimensions['B'].width = 60

# --- Instruções
wi = wb.create_sheet('Instruções')
texto = [
    ('Legenda', 'Células amarelas: preencher. Células cinzas: calculadas (não editar). Coluna Conferência (vermelho): '
                'inconsistências a resolver. Linha cinza em itálico no Roteiro: exemplo. Protocolo completo: '
                'docs/pesquisa/estudo2_tarefas.md.'),
    ('Condições', 'S = Chrome sem extensão (perfil S). C = Chrome com Accessify v1.1.0-jis_submission (perfil C), acionado '
                  'pelo painel lateral em cada página nova do site da tarefa, antes de procurar a informação: Alt+Shift+A '
                  '(abre o painel e inicia a análise), ouvir o resultado, "Aplicar Otimizações", ouvir a confirmação e voltar '
                  'à página. A sequência exata de teclas (provavelmente F6 para alternar painel e página) é fixada no piloto. '
                  'Não acionar duas vezes na mesma página (a segunda análise recarrega a página). Sites externos ao da tarefa '
                  'não recebem o Accessify.'),
    ('Tempos', 't0 = o(a) facilitador(a) diz "pode começar". t1 = a avaliadora diz a resposta final ("minha resposta é…") ou '
               '"desisto"; se ela corrigir a resposta, valem a última resposta e o seu t1. Releituras da instrução a pedido '
               'ficam dentro do tempo e são contadas em Repetições. Tempo total = t1 − t0. Tempo no painel (só em C) = soma '
               'dos intervalos entre Alt+Shift+A e a volta do foco à página (primeira fala do NVDA vinda da página depois da '
               'confirmação). Tempo líquido = total − painel.'),
    ('Teto', '300 s de tempo líquido nas duas condições. Ao vivo, o(a) facilitador(a) pausa o cronômetro durante o painel e '
             'encerra a tarefa aos 300 s; execução encerrada pelo(a) facilitador(a) ou com tempo líquido acima de 300 s '
             'conta como F.'),
    ('Resultado', 'S = resposta correta segundo o critério da linha, sem ajuda, dentro do teto. P = parcialmente correta, como '
                  'definido no critério. F = errada, desistência, ajuda ou teto. Ajuda = indicar caminho, conteúdo ou atalho. '
                  'Lembrar de acionar o Accessify não é ajuda: o(a) facilitador(a) lembra assim que ela começar a explorar uma '
                  'página nova sem o painel, e conta em "Lembretes".'),
    ('Falha da extensão', 'Se o painel mostrar erro ou não anunciar o fim da análise em 60 s, a avaliadora segue sem acionar '
                          'de novo; escrever "C degradada" em Observações (entra numa análise de sensibilidade).'),
    ('SEQ (Single Ease Question)', 'Logo após cada tarefa: "No geral, quão fácil ou difícil foi esta tarefa, de 1 (muito '
                                   'difícil) a 7 (muito fácil)?"'),
    ('Estratégia', 'Como a avaliadora chegou à resposta: cabeçalhos, lista de links, busca do NVDA, leitura linear, tabela, '
                   'menu do site, busca do site ou várias (texto livre permitido).'),
    ('Barreiras e falas do NVDA', 'Trecho literal do que o NVDA falou nos pontos de dificuldade (ex.: "link Abrir link 11", '
                                  '"gráfico"), com o tempo no vídeo. Alimenta o Barrier Walkthrough (impacto 1–3 × '
                                  'persistência).'),
    ('Antes de cada tarefa', 'O(a) facilitador(a) abre a URL inicial, espera a carga, cala o NVDA (Ctrl) e só então lê a ação; '
                             'a avaliadora não toca no teclado até "pode começar". Sem pensar em voz alta (comentários depois '
                             'da resposta), sem ensinar atalhos, sem baixar arquivos.'),
]
for i, (k, v) in enumerate(texto):
    a = wi.cell(row=1 + i, column=1, value=k)
    a.font = F_BOLD
    a.alignment = WRAP
    b = wi.cell(row=1 + i, column=2, value=v)
    b.font = F_BASE
    b.alignment = WRAP
wi.column_dimensions['A'].width = 24
wi.column_dimensions['B'].width = 120

wb._sheets = [wb[n] for n in ('Instruções', 'Sessão', 'Roteiro', 'Resumo', 'Tarefas')]
wb.active = 2
wb.save(SAIDA_XLSX)
print('formulário:', os.path.relpath(SAIDA_XLSX), '· linhas do Roteiro', primeira, '–', ultima)

# ------------------------------------------------------------------ markdown (cartões)
if SAIDA_MD:
    md = []
    for conj, titulo in (('Núcleo', 'núcleo'), ('Complementar', 'complementar'), ('Reserva', 'reservas')):
        md.append(f'### Cartões — {titulo}\n')
        for t in [x for x in T if x['conj'].startswith(conj)]:
            grupo = f' · grupo {t["grupo"]}' if t['grupo'] != '—' else ''
            md.append(f'#### {t["id"]} — {t["titulo"]}\n')
            md.append(f'{t["site"]}{grupo} · início: <{t["url"]}> · {t["pag"]} página(s) no caminho · '
                      f'estimativa {t["est_s"]} s sem extensão, {t["est_c"]} s com extensão\n')
            md.append('| Variante | Ação (ler em voz alta) | Saída esperada | Critério |')
            md.append('|---|---|---|---|')
            for v in ('a', 'b'):
                md.append(f'| {v.upper()} | {t["instr_" + v]} | {t["saida_" + v]} | {t["crit_" + v]} |')
            md.append('')
            md.append(f'- **Caminho esperado:** {t["caminho"]}')
            md.append(f'- **Efeito previsto da extensão:** {t["efeito"]}')
            md.append(f'- **Riscos e cuidados:** {t["riscos"]}\n')
    with open(SAIDA_MD, 'w') as f:
        f.write('\n'.join(md).strip() + '\n')
    print('cartões:', SAIDA_MD)
