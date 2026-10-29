#!/usr/bin/env python3
"""Estima t0 e t1 das tarefas da sessão de 28/10/2026 do Estudo 2 enquanto o vídeo não é
codificado, cruzando a transcrição automática do Meet (.docx, com uma marca de tempo a cada
5 min) com a transcrição do chat (.sbv, horário de cada mensagem). Uso:
   python3 -I eval/tools/estudo2_tempos.py TRANSCRICAO.docx CHAT.sbv [--xlsx FORMULARIO.xlsx]
Sem --xlsx, imprime a validação e a tabela em Markdown. Com --xlsx, grava no Roteiro t0 e t1
estimados (N, O), o tempo no painel das linhas C (S = 0: modo automático, sem Alt+Shift+A entre
t0 e t1) e a faixa de cada tempo (coluna AH). Depois, recalcular as fórmulas (LibreOffice).
O .sbv contém a chave do OpenRouter colada na sessão: o script só lê os horários das mensagens.

Hipóteses: vídeo, transcrição e chat têm o mesmo zero; a transcrição está em ordem cronológica;
a marca "00:05:00" fica entre a última fala que começou antes e a primeira que começou depois.
Definições do protocolo (docs/pesquisa/estudo2_tarefas.md), adaptadas porque na sessão não houve
"pode começar": t0 = fim da instrução falada, com as repetições do facilitador antes de a
avaliadora começar; t1 = fim da última resposta dela (mesmo depois de ajuda) ou da desistência;
sem nenhuma das duas, a fala do facilitador que encerrou a tarefa. As linhas vêm da codificação
(coluna AF do formulário), conferida por um segundo leitor.

Âncoras: as marcas de 5 min e as 32 mensagens do chat da sessão (27 URLs e títulos das tarefas,
mais atalhos, comandos de instalação e a chave), cada uma casada com a fala que a menciona. Fala no futuro ("vou colar") começa antes da mensagem
(<=); no passado ou presente ("colei", "tá no chat") começa depois (>=); "~" = simultânea.
Faixa de cada tempo: o que cabe entre as âncoras supondo fala rápida (FAST palavras/s) e nenhum
silêncio; vale sob as hipóteses acima. Estimativa: média de dois modelos com fala no ritmo
típico (TYP, mediana das janelas da entrevista final, conversa sem navegação):
  A  interpola cada linha a partir da âncora anterior e da seguinte, com peso inverso à
     distância em linhas;
  B  põe o silêncio que sobra entre duas âncoras dentro das tarefas (entre t0 e t1), onde a
     avaliadora navega calada (o NVDA dela não entra na transcrição).
A validação deixa cada mensagem do chat de fora e prevê o horário da fala que a menciona.
"""
import html
import os
import re
import statistics as st
import sys
import zipfile

FAST = 6.0       # palavras/s: fala rápida, para a faixa
MTOL = 3.0       # s: tolerância das marcas de 5 min
EPS = 0.05       # peso do silêncio fora das tarefas no modelo B
TETO = 300       # s líquidos (protocolo)
ENTREVISTA = ('3:05:00', '3:10:00', '3:15:00', '3:20:00')   # janelas que calibram TYP

# Mensagem do chat (início no .sbv) -> linha da transcrição que a menciona, relação, tolerância (s)
CHAT = [
    ('00:12:11.539', 139, '<=', 5, 'chrome://extensions/shortcuts'),
    ('00:22:26.273', 274, '<=', 5, 'comando taskkill'),
    ('00:22:43.371', 277, '>=', 5, 'comando setx OLLAMA_ORIGINS'),
    ('00:22:58.718', 279, '>=', 5, 'comando start ollama app'),
    ('00:35:30.422', 400, '>=', 5, 'chave do OpenRouter'),
    ('00:52:36.983', 633, '<=', 5, 'título SOL-1 A'),
    ('00:52:36.983', 634, '>=', 5, 'título SOL-1 A'),
    ('00:54:53.204', 648, '<=', 5, 'URL CAPES-1 A'),
    ('00:54:53.204', 651, '>=', 5, 'URL CAPES-1 A'),
    ('00:55:31.364', 660, '>=', 5, 'nome do periódico CAPES-1 A'),
    ('00:59:32.779', 690, '>=', 5, 'URL IHC-2 A'),
    ('01:05:18.670', 734, '>=', 5, 'URL CSBC-2 A'),
    ('01:15:07.080', 850, '<=', 5, 'URL SOL-2 A'),
    ('01:15:39.496', 851, '<=', 5, 'título SOL-2 A'),
    ('01:15:39.496', 858, '>=', 5, 'título SOL-2 A'),
    ('01:25:17.417', 953, '~', 20, 'URL CAPES-2 A'),
    ('01:33:34.158', 1063, '<=', 5, 'URL IHC-1 A'),
    ('01:45:58.724', 1163, '>=', 5, 'URL EDU-1 A'),
    ('01:51:32.057', 1236, '>=', 5, 'URL SOL-1 B'),
    ('01:51:57.066', 1237, '>=', 5, 'título SOL-1 B'),
    ('01:53:39.212', 1245, '<=', 5, 'URL CAPES-1 B'),
    ('01:53:54.855', 1249, '>=', 5, 'nome do periódico CAPES-1 B'),
    ('01:56:37.227', 1266, '~', 20, 'URL IHC-2 B'),
    ('01:59:13.395', 1290, '>=', 5, 'URL CSBC-2 B'),
    ('02:05:47.777', 1349, '>=', 5, 'URL SOL-2 B'),
    ('02:08:44.853', 1360, '>=', 5, 'URL CAPES-2 B'),
    ('02:10:45.305', 1392, '>=', 5, 'URL IHC-1 B'),
    ('02:12:26.713', 1394, '>=', 5, 'título IHC-1 B'),
    ('02:25:49.484', 1520, '>=', 5, 'URL EDU-1 B'),
    ('02:26:56.848', 1527, '~', 20, 'URL JIS-1 A'),
    ('02:27:39.982', 1529, '<=', 5, 'título JIS-1 A'),
    ('02:37:37.676', 1594, '>=', 5, 'URL EDU-3 A'),
    ('02:39:35.933', 1621, '~', 20, 'URL SBC (SBC-1 A e SBC-2 A)'),
    ('02:54:53.557', 1710, '>=', 5, 'URL JIS-1 B'),
    ('02:57:54.728', 1747, '>=', 5, 'URL EDU-3 B'),
]

# Linha do Roteiro, tarefa, condição, linha de t0, linha de t1, como terminou
TAREFAS = [
    (6, 'SOL-1 A', 'S', 638, 646, 'resposta'),
    (7, 'CAPES-1 A', 'S', 661, 683, 'resposta corrigida'),
    (8, 'IHC-2 A', 'S', 702, 719, 'resposta'),
    (9, 'CSBC-2 A', 'S', 737, 744, 'resposta'),
    (12, 'SOL-2 A', 'C', 861, 880, 'resposta'),
    (13, 'CAPES-2 A', 'C', 954, 971, 'resposta'),
    (14, 'IHC-1 A', 'C', 1082, 1161, 'encerrada'),
    (15, 'EDU-1 A', 'C', 1211, 1219, 'resposta'),
    (17, 'SOL-1 B', 'C', 1239, 1244, 'resposta'),
    (18, 'CAPES-1 B', 'C', 1250, 1254, 'resposta'),
    (19, 'IHC-2 B', 'C', 1274, 1276, 'resposta'),
    (20, 'CSBC-2 B', 'C', 1291, 1302, 'última fala, "não tem como"; encerrada em L1309'),
    (22, 'SOL-2 B', 'S', 1354, 1358, 'resposta'),
    (23, 'CAPES-2 B', 'S', 1363, 1367, 'resposta'),
    (24, 'IHC-1 B', 'S', 1395, 1482, 'resposta, "não tem"; encerrada em L1485'),
    (25, 'EDU-1 B', 'S', 1521, 1522, 'resposta'),
    (27, 'JIS-1 A', 'S', 1531, 1545, 'resposta; depois, ajuda e encerramento em L1566'),
    (28, 'EDU-3 A', 'S', 1608, 1613, 'resposta'),
    (30, 'SBC-1 A', 'C', 1639, 1651, 'resposta'),
    (31, 'SBC-2 A', 'C', 1658, 1689, 'desistência'),
    (33, 'JIS-1 B', 'C', 1718, 1746, 'encerrada'),
    (34, 'EDU-3 B', 'C', 1752, 1756, 'resposta'),
    (36, 'SBC-1 B', 'S', 1759, 1774, 'resposta'),
    (37, 'SBC-2 B', 'S', 1776, 1791, 'encerrada'),
]


def seg(s):
    p = [float(x) for x in s.split(':')]
    return p[0] * 3600 + p[1] * 60 + p[2]


def hms(x):
    x = int(round(x))
    return f'{x // 3600}:{x % 3600 // 60:02d}:{x % 60:02d}'


# ------------------------------------------------------------------ entrada
def le_docx(caminho):
    """Parágrafos não vazios do .docx, na ordem; a linha n da transcrição é o n-ésimo."""
    xml = zipfile.ZipFile(caminho).read('word/document.xml').decode('utf-8')
    out = []
    for p in re.findall(r'<w:p[ >].*?</w:p>', xml, flags=re.S):
        t = html.unescape(''.join(re.findall(r'<w:t(?: [^>]*)?>(.*?)</w:t>', p, flags=re.S)))
        if t.strip():
            out.append(t)
    return out


def le_sbv_horarios(caminho):
    """Só os horários de início das mensagens (o texto não é lido)."""
    return {m.group(1) for m in re.finditer(r'^(\d{2}:\d{2}:\d{2}\.\d{3}),', open(caminho, encoding='utf-8').read(), flags=re.M)}


def carrega(docx, sbv):
    global FALAS, MARCAS, N
    FALAS, MARCAS = {}, {}
    for n, t in enumerate(le_docx(docx), 1):
        if re.fullmatch(r'\d{2}:\d{2}:\d{2}', t):
            MARCAS[n] = seg(t)
        elif n > 4:   # 1–4: cabeçalho (início da reunião, participantes)
            fala = t.partition(': ')[2] or t
            FALAS[n] = max(1, len(fala.split()))
    N = max(FALAS)
    horarios = le_sbv_horarios(sbv)
    faltam = [c[0] for c in CHAT if c[0] not in horarios]
    if faltam:
        sys.exit(f'mensagens do chat não encontradas no .sbv: {faltam}')


# ------------------------------------------------------------------ restrições
def dur(i, taxa):
    return FALAS.get(i, 0) / taxa


def fala(a, b, taxa):
    """Fala das linhas a..b-1."""
    return sum(dur(i, taxa) for i in range(a, b))


def prox(n):
    return next(k for k in range(n + 1, N + 2) if k in FALAS or k == N + 1)


def ant(n):
    return next(k for k in range(n - 1, 0, -1) if k in FALAS)


def restricoes(fora=None):
    """Listas (linha, tempo) com início(linha) >= tempo e com início(linha) <= tempo."""
    lo, hi = [], []
    for m, T in MARCAS.items():
        lo.append((prox(m), T - MTOL))
        hi.append((ant(m), T + MTOL))
    for i, (t, n, op, tol, _) in enumerate(CHAT):
        if i == fora:
            continue
        if op in ('>=', '~'):
            lo.append((n, seg(t) - tol))
        if op in ('<=', '~'):
            hi.append((n, seg(t) + tol))
    return lo, hi


def faixa_inicio(L, lo, hi):
    a = max(T + fala(k, L, FAST) for k, T in lo if k <= L)
    b = min(T - fala(L, k, FAST) for k, T in hi if k >= L)
    return a, b


def faixa_fim(L, lo, hi):
    """Fim da fala L: depois do início mais cedo mais a fala rápida, antes do início da seguinte."""
    return faixa_inicio(L, lo, hi)[0] + dur(L, FAST), faixa_inicio(prox(L), lo, hi)[1]


# ------------------------------------------------------------------ modelo A
def inicio_a(L, lo, hi):
    f = max(((k, T) for k, T in lo if k <= L), default=None)
    g = min(((k, T) for k, T in hi if k >= L), default=None)
    ef = f[1] + fala(f[0], L, TYP) if f else None
    eg = g[1] - fala(L, g[0], TYP) if g else None
    if ef is None or eg is None:
        return ef if eg is None else eg
    df, dg = L - f[0] + 1, g[0] - L + 1
    return (ef / df + eg / dg) / (1 / df + 1 / dg)


# ------------------------------------------------------------------ modelo B
NAS_TAREFAS = {i for *_, a, b, _f in TAREFAS for i in range(a, b)}   # silêncio depois da linha i


def ancoras_b(fora=None):
    A = [(prox(m), T, 'm') for m, T in MARCAS.items()]
    A += [(n, seg(t), op) for i, (t, n, op, _, _) in enumerate(CHAT) if i != fora]
    A.sort()
    out = []
    for a in A:   # descarta a âncora do chat que contradiz a anterior; marca tem prioridade
        if out and a[1] < out[-1][1] + fala(out[-1][0], a[0], FAST):
            if a[2] != 'm':
                continue
            out.pop()
        if out and out[-1][0] == a[0]:
            continue
        out.append(a)
    return out


def inicios_b(A):
    S = {}
    for (k1, T1, _), (k2, T2, _) in zip(A, A[1:]):
        ids = [i for i in range(k1, k2) if i in FALAS]
        f = sum(dur(i, TYP) for i in ids)
        esc = min(1.0, (T2 - T1) / f) if f else 1.0
        folga = max(0.0, T2 - T1 - f)
        w = [1.0 if i in NAS_TAREFAS else EPS for i in ids]
        t = T1
        for i, wi in zip(ids, w):
            S[i] = t
            t += dur(i, TYP) * esc + folga * wi / sum(w)
    return S


# ------------------------------------------------------------------ cálculo
def calcula():
    global TYP
    mk = sorted(MARCAS.items())
    TYP = st.median(sum(FALAS.get(i, 0) for i in range(m1 + 1, m2)) / (T2 - T1)
                    for (m1, T1), (m2, T2) in zip(mk, mk[1:]) if hms(T1) in ENTREVISTA)
    lo, hi = restricoes()
    # validação: cada mensagem do chat de fora
    val = []
    for i, (t, n, op, tol, rot) in enumerate(CHAT):
        lo2, hi2 = restricoes(fora=i)
        a, b = faixa_inicio(n, lo2, hi2)
        ea = inicio_a(n, lo2, hi2)
        eb = inicios_b(ancoras_b(fora=i)).get(n)
        T = seg(t)
        val.append(dict(rot=rot, linha=n, op=op, chat=t[:8], erro_a=ea - T, erro_b=(eb - T) if eb is not None else None,
                        dentro=a - tol <= T <= b + tol))
    SB = inicios_b(ancoras_b())
    res = []
    for linha, tarefa, cond, l0, l1, fim in TAREFAS:
        f0, f1 = faixa_fim(l0, lo, hi), faixa_fim(l1, lo, hi)
        a0 = inicio_a(l0, lo, hi) + dur(l0, TYP)
        a1 = inicio_a(l1, lo, hi) + dur(l1, TYP)
        b0 = SB[l0] + dur(l0, TYP)
        b1 = SB[l1] + dur(l1, TYP)
        entre = fala(prox(l0), l1, FAST) + dur(l1, FAST)
        t0 = min(max((a0 + b0) / 2, f0[0]), f0[1])
        t1 = min(max((a1 + b1) / 2, f1[0]), f1[1], 10 ** 9)
        t1 = max(t1, t0 + entre)
        t0, t1 = round(t0), round(t1)   # o total bate com o do formulário (t1 − t0 em segundos inteiros)
        dmin, dmax = max(entre, f1[0] - f0[1]), f1[1] - f0[0]
        teto = 'abaixo (garantido)' if dmax <= TETO else 'acima (garantido)' if dmin > TETO else 'indeterminado'
        res.append(dict(linha=linha, tarefa=tarefa, cond=cond, l0=l0, l1=l1, fim=fim, t0=t0, t1=t1, f0=f0, f1=f1,
                        total=t1 - t0, dmin=dmin, dmax=dmax, da=a1 - a0, db=b1 - b0, teto=teto))
    return val, res


def resumo_validacao(val):
    ea = [abs(v['erro_a']) for v in val]
    eb = [abs(v['erro_b']) for v in val if v['erro_b'] is not None]
    em = [abs((v['erro_a'] + v['erro_b']) / 2) for v in val if v['erro_b'] is not None]
    return dict(n=len(val), dentro=sum(v['dentro'] for v in val), a=(st.median(ea), max(ea)), b=(st.median(eb), max(eb)),
                media=(st.median(em), max(em)))


def texto_ah(r):
    return (f"Estimado (transcrição + chat; conferir no vídeo). t0 ≈ {hms(r['t0'])} (L{r['l0']:04d}; faixa {hms(r['f0'][0])}–{hms(r['f0'][1])}); "
            f"t1 ≈ {hms(r['t1'])} (L{r['l1']:04d}, {r['fim']}; faixa {hms(r['f1'][0])}–{hms(r['f1'][1])}). "
            f"Total ≈ {r['total']:.0f} s (faixa {r['dmin']:.0f}–{r['dmax']:.0f} s; modelos A/B: {r['da']:.0f}/{r['db']:.0f} s); "
            f"teto de 300 s: {r['teto']}." + (' Tempo no painel = 0: modo automático, sem Alt+Shift+A entre t0 e t1.' if r['cond'] == 'C' else ''))


def grava_xlsx(caminho, res):
    try:
        from copy import copy
        from openpyxl import load_workbook
    except ImportError:
        sys.exit('requer openpyxl (pip install openpyxl)')
    wb = load_workbook(caminho)
    ws = wb['Roteiro']
    ws['AH3'] = 'Tempos estimados (transcrição + chat): faixa e âncoras'
    for c in ('font', 'fill', 'border', 'alignment'):
        setattr(ws['AH3'], c, copy(getattr(ws['AF3'], c)))
    ws.column_dimensions['AH'].width = ws.column_dimensions['AF'].width or 40
    for r in res:
        n = r['linha']
        if ws[f'D{n}'].value != r['tarefa'].split()[0] or ws[f'E{n}'].value != r['tarefa'].split()[1]:
            sys.exit(f'linha {n} do Roteiro não é {r["tarefa"]}')
        ws[f'N{n}'], ws[f'O{n}'] = hms(r['t0']), hms(r['t1'])
        if r['cond'] == 'C':
            ws[f'S{n}'] = 0
        ws[f'AH{n}'] = texto_ah(r)
        for c in ('font', 'fill', 'border', 'alignment'):
            setattr(ws[f'AH{n}'], c, copy(getattr(ws[f'AF{n}'], c)))
    wb.save(caminho)


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    carrega(sys.argv[1], sys.argv[2])
    val, res = calcula()
    v = resumo_validacao(val)
    print(f'TYP = {TYP:.2f} palavras/s; mensagens do chat dentro da faixa prevista sem elas: {v["dentro"]}/{v["n"]}')
    print(f'erro |previsto − chat| (mediana, máx.): A {v["a"][0]:.0f}/{v["a"][1]:.0f} s; B {v["b"][0]:.0f}/{v["b"][1]:.0f} s; '
          f'média dos dois {v["media"][0]:.0f}/{v["media"][1]:.0f} s')
    print()
    print('| Linha | Tarefa | Cond. | t0 ≈ | t1 ≈ | Total ≈ (s) | Faixa do total (s) | A / B (s) | Teto de 300 s |')
    print('|---|---|---|---|---|---|---|---|---|')
    for r in res:
        print(f"| {r['linha']} | {r['tarefa']} | {r['cond']} | {hms(r['t0'])} | {hms(r['t1'])} | {r['total']:.0f} | "
              f"{r['dmin']:.0f}–{r['dmax']:.0f} | {r['da']:.0f} / {r['db']:.0f} | {r['teto']} |")
    if '--xlsx' in sys.argv:
        grava_xlsx(sys.argv[sys.argv.index('--xlsx') + 1], res)
        print(f'\ngravado: {sys.argv[sys.argv.index("--xlsx") + 1]} (recalcular as fórmulas antes de abrir)')


if __name__ == '__main__':
    main()
